"""Genera el informe editable y sus diagramas a partir del avance comprobado."""
import json
import math
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "entrega"
DIAGRAMAS = RAIZ / "docs" / "diagramas"
SALIDA.mkdir(exist_ok=True)
DIAGRAMAS.mkdir(exist_ok=True)
REGISTRO = json.loads((RAIZ / "docs/evidencias/resultados.json").read_text(encoding="utf-8"))
FUENTE = "C:/Windows/Fonts/arial.ttf"
NEGRITA = "C:/Windows/Fonts/arialbd.ttf"


def diagrama(nombre, ancho, alto):
    imagen = Image.new("RGB", (ancho, alto), "white")
    return imagen, ImageDraw.Draw(imagen)


def caja(d, xy, titulo, lineas=()):
    x, y, w, h = xy
    d.rectangle((x, y, x + w, y + h), fill="#f6f8fa", outline="#354354", width=3)
    fuente = ImageFont.truetype(NEGRITA, 27)
    d.text((x + w / 2, y + 14), titulo, fill="#111111", font=fuente, anchor="mt")
    d.line((x, y + 55, x + w, y + 55), fill="#354354", width=2)
    for i, linea in enumerate(lineas):
        d.text((x + 16, y + 67 + i * 33), linea, fill="#222222", font=ImageFont.truetype(FUENTE, 25))


def relacion(d, puntos, tipo="flecha", etiqueta=None, pos=None, discontinua=False):
    for a, b in zip(puntos, puntos[1:]):
        if discontinua:
            largo = math.dist(a, b)
            for i in range(0, int(largo), 18):
                fin = min(i + 10, largo)
                d.line((a[0] + (b[0] - a[0]) * i / largo, a[1] + (b[1] - a[1]) * i / largo,
                        a[0] + (b[0] - a[0]) * fin / largo, a[1] + (b[1] - a[1]) * fin / largo), fill="#354354", width=3)
        else:
            d.line((a, b), fill="#354354", width=3)
    a, b = puntos[-2:]
    ang = math.atan2(b[1] - a[1], b[0] - a[0])
    if tipo in ("flecha", "herencia"):
        izq = (b[0] - 22 * math.cos(ang) + 11 * math.sin(ang), b[1] - 22 * math.sin(ang) - 11 * math.cos(ang))
        der = (b[0] - 22 * math.cos(ang) - 11 * math.sin(ang), b[1] - 22 * math.sin(ang) + 11 * math.cos(ang))
        if tipo == "herencia":
            d.polygon([b, izq, der], fill="white", outline="#354354", width=3)
        else:
            d.line((izq, b, der), fill="#354354", width=3)
    if tipo == "composicion":
        a, b = puntos[:2]
        ang = math.atan2(b[1] - a[1], b[0] - a[0])
        puntos_rombo = [a,
            (a[0]+14*math.cos(ang)+9*math.sin(ang),a[1]+14*math.sin(ang)-9*math.cos(ang)),
            (a[0]+28*math.cos(ang),a[1]+28*math.sin(ang)),
            (a[0]+14*math.cos(ang)-9*math.sin(ang),a[1]+14*math.sin(ang)+9*math.cos(ang))]
        d.polygon(puntos_rombo, fill="#354354")
    if etiqueta:
        d.text(pos, etiqueta, font=ImageFont.truetype(FUENTE, 24), fill="#222222")


def crear_diagramas():
    im, d = diagrama("dominio", 1500, 1120)
    caja(d, (550, 20, 380, 125), "Juego", ["+ iniciar_partida()"])
    caja(d, (550, 230, 380, 155), "Nivel", ["+ actualizar(dt, teclas)", "+ resolver_interacciones()"])
    caja(d, (550, 490, 380, 155), "Personaje  {abstracta}", ["- _x, _y, _velocidad", "+ update() {abstracto}"])
    caja(d, (140, 760, 420, 145), "Prisionero", ["+ update()", "+ recoger_llave()"])
    caja(d, (950, 760, 420, 145), "Guardia", ["+ update()", "+ puede_detectar(prisionero)"])
    caja(d, (140, 980, 420, 115), "HabilidadCarrera", ["+ activar() / actualizar(dt)"])
    caja(d, (25, 230, 375, 115), "Camara", ["+ detectar(prisionero, muros)"])
    caja(d, (1070, 160, 400, 115), "Llave", ["+ recoger(prisionero)"])
    caja(d, (1070, 390, 400, 115), "PuertaSalida", ["+ puede_salir(prisionero)"])
    relacion(d, [(740,145),(740,230)],"composicion","1        1",(762,175))
    relacion(d, [(740,385),(740,490)],"composicion","1        2",(762,424))
    relacion(d, [(550,285),(400,285)],"composicion","1    1",(428,250))
    relacion(d, [(930,270),(1000,270),(1000,215),(1070,215)],"composicion","1   1",(959,165))
    relacion(d, [(930,340),(1000,340),(1000,450),(1070,450)],"composicion","1   1",(1020,462))
    relacion(d, [(350,760),(350,690),(665,690),(665,645)],"herencia")
    relacion(d, [(1160,760),(1160,690),(805,690),(805,645)],"herencia")
    relacion(d, [(350,905),(350,980)],"composicion","1        1",(375,922))
    d.text((600,1000), "Rombo lleno: composición\nTriángulo vacío: herencia\nMultiplicidades del prototipo", fill="#222222", font=ImageFont.truetype(FUENTE,26), spacing=12)
    im.save(DIAGRAMAS / "01_dominio.png")

    im, d = diagrama("state", 1500, 740)
    caja(d,(25,35,350,130),"Juego",["+ cambiar_estado(estado)"])
    caja(d,(720,35,580,185),"EstadoJuego  {abstracta}",["+ manejar_entrada(tecla)","+ actualizar(dt, teclas)"])
    relacion(d,[(375,93),(720,93)],"flecha","1       delega       1",(418,43))
    for x, nombre in [(25,"Jugando"),(400,"Pausa"),(775,"Derrota"),(1150,"Victoria")]:
        caja(d,(x,355,320,130),nombre,["Entrada y actualización"])
        relacion(d,[(x+160,355),(x+160,285),(1010,285),(1010,220)],"herencia")
    d.text((35,540),"Jugando -> Pausa: P o ESC     |     Pausa -> Jugando: P o ESC",font=ImageFont.truetype(FUENTE,29),fill="#222222")
    d.text((35,594),"Jugando -> Derrota: detección     |     Jugando -> Victoria: salida con llave",font=ImageFont.truetype(FUENTE,29),fill="#222222")
    d.text((35,648),"Derrota / Victoria -> Jugando: R, con un nivel nuevo y puntaje cero",font=ImageFont.truetype(FUENTE,29),fill="#222222")
    im.save(DIAGRAMAS / "02_state.png")

    im,d=diagrama("observer",1500,760)
    caja(d,(25,60,460,245),"Puntaje",["- _valor", "- _observadores", "+ suscribir(observador)","+ sumar(cantidad)","- _notificar()"])
    caja(d,(905,60,560,175),"ObservadorPuntaje  {abstracta}",["+ actualizar_puntaje(valor)","  {abstracto}"])
    caja(d,(905,395,560,145),"PanelPuntaje",["+ valor", "+ actualizar_puntaje(valor)"])
    relacion(d,[(485,145),(905,145)],"flecha","1    notifica    0..*",(548,91))
    relacion(d,[(1180,395),(1180,235)],"herencia")
    d.text((45,390),"Llave recogida",font=ImageFont.truetype(NEGRITA,31),fill="#222222")
    d.text((45,445),"sumar(100) -> _notificar()",font=ImageFont.truetype(FUENTE,29),fill="#222222")
    d.text((45,590),"El panel se suscribe una vez. Reiniciar también notifica el valor cero.",font=ImageFont.truetype(FUENTE,30),fill="#222222")
    d.text((45,650),"La vista dibuja PanelPuntaje.valor; Puntaje no conoce Pygame.",font=ImageFont.truetype(FUENTE,30),fill="#222222")
    im.save(DIAGRAMAS / "03_observer.png")


crear_diagramas()
doc = Document()
seccion = doc.sections[0]
seccion.page_width = Cm(21)
seccion.page_height = Cm(29.7)
seccion.top_margin = Cm(1.8)
seccion.bottom_margin = Cm(1.8)
seccion.left_margin = Cm(2)
seccion.right_margin = Cm(2)
for nombre in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
    estilo = doc.styles[nombre]
    estilo.font.name = "Calibri"
    estilo.font.color.rgb = RGBColor(0,0,0)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Normal"].paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(28)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
doc.styles["Caption"].font.size = Pt(9)
footer = seccion.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Fuga de la Prisión  |  ").font.size = Pt(9)
campo = OxmlElement("w:fldSimple")
campo.set(qn("w:instr"), "PAGE")
footer._p.append(campo)
md=[]


def h(t,n=1):
    doc.add_heading(t,n)
    md.append("#"*(n+1)+" "+t+"\n")


def p(t):
    doc.add_paragraph(t)
    md.append(t+"\n")


def tabla(cabecera, filas, anchos):
    t=doc.add_table(rows=1,cols=len(cabecera))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,anchos):
        c.width=Cm(w)
    for i,v in enumerate(cabecera):
        t.rows[0].cells[i].text=v
    for fila in filas:
        celdas=t.add_row().cells
        for i,v in enumerate(fila):
            celdas[i].text=str(v)
    for r,fila in enumerate(t.rows):
        pr=fila._tr.get_or_add_trPr()
        evitar=OxmlElement("w:cantSplit")
        pr.append(evitar)
        if r==0:
            repetir=OxmlElement("w:tblHeader")
            pr.append(repetir)
        for c,celda in enumerate(fila.cells):
            celda.width=Cm(anchos[c])
            celda.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=celda._tc.get_or_add_tcPr()
            margenes=OxmlElement("w:tcMar")
            for lado in ("top","left","bottom","right"):
                n=OxmlElement("w:"+lado)
                n.set(qn("w:w"),"80")
                n.set(qn("w:type"),"dxa")
                margenes.append(n)
            cp.append(margenes)
            bordes=OxmlElement("w:tcBorders")
            for lado in ("top","left","bottom","right"):
                n=OxmlElement("w:"+lado)
                n.set(qn("w:val"),"single")
                n.set(qn("w:sz"),"4")
                n.set(qn("w:color"),"D9D9D9")
                bordes.append(n)
            cp.append(bordes)
            shade=OxmlElement("w:shd")
            shade.set(qn("w:fill"),"DCE6EE" if r==0 else ("F5F7F9" if r%2==0 else "FFFFFF"))
            cp.append(shade)
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_after=Pt(2)
                parrafo.paragraph_format.line_spacing=1.0
                for run in parrafo.runs:
                    run.font.size=Pt(9)
                    run.bold=(r==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    md.append("| "+" | ".join(cabecera)+" |")
    md.append("| "+" | ".join(["---"]*len(cabecera))+" |")
    for fila in filas:
        md.append("| "+" | ".join(str(x) for x in fila)+" |")
    md.append("")


def imagen(ruta, ancho, titulo):
    parrafo=doc.add_paragraph()
    parrafo.alignment=WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.keep_with_next=True
    dibujo=parrafo.add_run().add_picture(str(ruta),width=Inches(ancho))
    dibujo._inline.docPr.set("descr",titulo)
    doc.add_paragraph(titulo,"Caption")
    md.append(f"![{titulo}]({ruta.relative_to(RAIZ / 'docs').as_posix()})\n")


def codigo(t):
    parrafo=doc.add_paragraph()
    parrafo.paragraph_format.space_after=Pt(9)
    run=parrafo.add_run(t)
    run.font.name="Consolas"
    run.font.size=Pt(10)
    md.append("```\n"+t+"\n```\n")


def pagina():
    doc.add_page_break()


doc.add_paragraph("Universidad Peruana Cayetano Heredia")
doc.add_paragraph("Programación Orientada a Objetos  |  Ciclo 2026-2")
doc.add_paragraph("Fuga de la Prisión", "Title")
doc.add_paragraph("Informe de avance parcial", "Subtitle")
doc.add_paragraph("Videojuego 2D desarrollado en Python")
doc.add_paragraph("Entrega programada para el 16 de octubre de 2026")
h("Integrantes",2)
for nombre in ["JOSE FERNANDO CHAVEZ LOZANO","MICKELE CARLOS TORRES VELAZCO","RICHARD AARON SORIANO CORDOVA","PIERO CEMEI ROMERO MEZA"]:
    p(nombre)
h("Resumen del avance",2)
p("El avance presenta un prototipo jugable de escape con un nivel, movimiento con colisiones, patrulla de un guardia, una cámara de vigilancia, recolección de una llave, salida bloqueada, puntuación y una habilidad de carrera. El jugador pierde al ser detectado y completa el pabellón al alcanzar la salida con la llave.")
p("La implementación utiliza una jerarquía de personajes actualizada polimórficamente, relaciones de composición y dependencia, y los patrones State y Observer. Se ejecutaron 20 pruebas automatizadas satisfactorias y una ruta completa hasta la victoria. Este informe relaciona los requisitos del curso con el código, documenta la validación y delimita las ampliaciones para el proyecto final.")
p("Versión parcial 0.1. Corte de implementación y pruebas: 8 de octubre de 2026. La fecha de entrega es posterior al corte y permite reservar tiempo para revisión y ensayo del equipo.")

pagina(); h("1 Objetivos y alcance")
h("Objetivo general",2)
p("Diseñar e implementar un videojuego 2D de escape de una prisión en Python que demuestre los principios de POO y al menos dos patrones de diseño indicados en la consigna del curso.")
h("Objetivos específicos",2)
p("Modelar personajes y objetos mediante UML; implementar movimiento, vigilancia, recolección y escape; aplicar herencia, polimorfismo, encapsulamiento, abstracción, composición y dependencia; separar el flujo de la partida y la actualización del puntaje; registrar avances en Git y comprobar las reglas con pruebas reproducibles.")
h("Alcance del parcial",2)
p("Se implementó el pabellón inicial con un prisionero, un guardia, una cámara, una llave y una salida. Incluye pausa, derrota, victoria, reinicio y carrera temporal. Los gráficos son figuras generadas por el código; no se emplearon imágenes o sonidos de terceros.")
tabla(["Requisito del curso [1]","Estado en el parcial"],[
    ("Videojuego 2D y puntuación, p. 1","Implementados con Pygame y marcador observable."),
    ("Distintos niveles, p. 1","Un nivel implementado; ampliación multinivel pendiente para el final."),
    ("Protagonista y personajes secundarios, p. 1","Prisionero y un guardia; más secundarios en niveles posteriores."),
    ("Tema 13 y seguridad, p. 2","Guardia, cámara y salida que exige llave."),
    ("UML y jerarquía con update(), p. 2","Diagramas de este informe y llamada común en Nivel."),
    ("Encapsulamiento, abstracción y relaciones, p. 2","Validados en clases y pruebas."),
    ("Git durante el desarrollo, p. 2","Historial local desde la planificación."),
    ("Al menos dos patrones, pp. 2 y 3","State y Observer implementados."),
    ("API externa, p. 3","No implementada. Obligatoriedad por confirmar: el título indica sugerencia y el cuerpo emplea debe.")
],[8,9])
p("La consigna describe el proyecto final y no establece una lista específica de entregables del parcial. El alcance presentado permite demostrar el diseño y un ciclo completo de juego; no implica que se hayan completado los requisitos del producto final.")

pagina(); h("2 Diseño del juego")
p("El jugador recorre un mapa con vista superior. Debe observar los campos de visión, recoger la llave y llegar a la salida. Los muros actúan como obstáculos físicos y visuales. El mapa inicial permite una ruta segura y también rutas expuestas para demostrar detección y derrota.")
imagen(RAIZ/"docs/evidencias/01_inicio.png",5.6,"Figura 1. Pantalla inicial renderizada por el prototipo. El prisionero es verde, el guardia rojo y la cámara amarilla.")
tabla(["Regla","Comportamiento implementado"],[
    ("Movimiento","WASD o flechas; misma velocidad total en diagonal; colisión circular contra muros."),
    ("Detección","Cono con alcance y bloqueo por muros; contacto con guardia también captura. Se evalúa el centro del jugador."),
    ("Llave y salida","Una llave aporta 100 puntos. Salir con ella aporta 500, una sola vez."),
    ("Prioridad","La captura prevalece sobre recoger la llave o alcanzar la salida en la misma actualización."),
    ("Carrera","Shift: velocidad por 1,5 durante 2 s; recarga de 5 s después del efecto."),
    ("Pausa y reinicio","P o Escape congelan la simulación; R tras victoria o derrota crea un nivel nuevo y pone el puntaje en cero.")
],[3.5,13.5])

pagina(); h("3 Requisitos funcionales y casos de uso")
tabla(["ID","Función","Criterio de aceptación"],[
    ("RF01","Iniciar","Mapa con todos los elementos y puntaje cero."),
    ("RF02","Mover","Controles responden y el personaje no atraviesa muros."),
    ("RF03","Patrullar","Guardia recorre dos puntos y orienta su campo visual."),
    ("RF04","Detectar","Guardia o cámara provocan derrota; los muros bloquean la visión."),
    ("RF05","Recoger","Llave se retira y suma 100 puntos exactamente una vez."),
    ("RF06","Escapar","Salida exige llave y termina el nivel con 600 puntos totales."),
    ("RF07","Notificar puntaje","El panel recibe los cambios mediante Observer."),
    ("RF08","Pausar","No avanzan posiciones, tiempo ni habilidad durante la pausa."),
    ("RF09","Reiniciar","Se restauran posiciones, llave, habilidad y puntaje."),
    ("RF10","Correr","Duración y recarga limitan la activación; se conservan colisiones.")
],[1.4,3.7,11.9])
h("CU01 Completar el pabellón",2)
p("Actor: jugador. Precondición: partida en estado Jugando. Flujo: el jugador avanza evitando vigilancia, recoge la llave, observa el aumento de 100 puntos y alcanza la salida. El nivel concede 500 puntos adicionales y Juego cambia a Victoria. Postcondición: simulación detenida, 600 puntos y opción de reinicio.")
p("Alternativas: si alcanza la salida sin llave, la partida continúa; si entra al campo de visión o contacta al guardia, cambia a Derrota; si pulsa P o Escape, pasa a Pausa y conserva su progreso.")
h("CU02 Reiniciar tras un resultado",2)
p("Actor: jugador. Precondición: Victoria o Derrota. Al pulsar R, Juego reinicia el objeto Puntaje, crea un nuevo Nivel y activa Jugando. Se restaura la llave, la posición inicial y la disponibilidad de la habilidad. Reiniciar durante la partida no es un control habilitado.")
h("Condiciones de calidad",2)
p("La lógica se puede probar sin abrir una ventana. La interfaz comunica objetivo, controles, puntaje, tiempo y disponibilidad de carrera. El bucle solicita 60 actualizaciones por segundo y limita a 0,05 s el paso de simulación ante demoras; esto no constituye una medición de rendimiento en otros equipos.")

pagina(); h("4 Modelo de clases del prototipo")
imagen(DIAGRAMAS/"01_dominio.png",6.55,"Figura 2. UML del dominio. Se muestran las multiplicidades de este prototipo de un nivel; las clases de patrones se detallan en las figuras siguientes.")
p("Juego coordina el estado, el nivel y la puntuación. Nivel construye sus personajes y objetos; por ello, su relación con esos elementos es de composición. Contiene exactamente dos personajes: un Prisionero y un Guardia. Al reiniciar se crea otra instancia de Nivel.")
p("Personaje declara update() como operación abstracta. Prisionero interpreta los controles y actualiza su carrera; Guardia sigue su ruta y orienta su visión. Ambos reciben una referencia al nivel, donde encuentran el tiempo de actualización y el contexto necesario.")
codigo("for personaje in self.personajes:\n    personaje.update()")
p("El fragmento pertenece a Nivel.actualizar(). La llamada es idéntica para ambos objetos y el comportamiento depende de su clase concreta. No hay una selección por tipo para decidir cuál update ejecutar.")
p("PuertaSalida.puede_salir(prisionero) depende temporalmente del objeto recibido y no lo conserva como atributo. La cámara se modela como dispositivo, no como subclase de Personaje. Llave se responsabiliza de su disponibilidad y recolección.")

pagina(); h("5 Aplicación del patrón State")
p("Juego conserva el estado activo y delega manejar_entrada() y actualizar(). Cada estado implementa ambas operaciones. Esta organización permite que pausa, derrota y victoria detengan la simulación sin incorporar esas decisiones en todos los objetos del mapa.")
imagen(DIAGRAMAS/"02_state.png",6.55,"Figura 3. Participantes de State y transiciones implementadas. El triángulo vacío representa herencia de la clase abstracta.")
tabla(["Estado","Entrada y actualización"],[
    ("Jugando","Actualiza Nivel; P o Escape activan Pausa y Shift solicita carrera. El resultado del nivel activa Victoria o Derrota."),
    ("Pausa","P o Escape reanudan. actualizar() no modifica la simulación."),
    ("Derrota","R inicia una nueva partida. La simulación permanece detenida."),
    ("Victoria","R inicia una nueva partida. La simulación permanece detenida.")
],[3,14])
codigo("def actualizar(self, dt, teclas):\n    resultado = self.juego.nivel.actualizar(dt, teclas)\n    if resultado == \"derrota\":\n        self.juego.cambiar_estado(Derrota(self.juego))\n    elif resultado == \"victoria\":\n        self.juego.cambiar_estado(Victoria(self.juego))")
p("Fragmento de Jugando en src/estados.py. Las pruebas comprueban que pausar conserva posiciones, tiempo y habilidad, y que reanudar restituye el movimiento. Perder el foco de la ventana también solicita pausa para evitar una captura mientras el jugador usa otra aplicación.")

pagina(); h("6 Observer y protección de los datos")
imagen(DIAGRAMAS/"03_observer.png",6.4,"Figura 4. Observer implementado mediante una clase abstracta. PanelPuntaje hereda el contrato de ObservadorPuntaje.")
p("Puntaje es el sujeto observado. Mantiene su valor y una lista de objetos suscritos. Al sumar o reiniciar, invoca actualizar_puntaje(valor) en cada observador. PanelPuntaje recibe el valor y Vista lo dibuja; el cálculo del puntaje no importa Pygame ni decide cómo se representa el texto.")
p("La suscripción del panel ocurre una sola vez al construir Juego. El mismo panel permanece entre reinicios. suscribir() evita duplicados y desuscribir() permite retirar un observador. Se comprueban el orden de notificaciones, la ausencia de duplicados y la desuscripción.")
h("Encapsulamiento y abstracción",2)
p("Puntaje expone sumar() y reiniciar() y rechaza incrementos negativos, cero, booleanos, texto o números no enteros. Personaje controla sus coordenadas mediante mover(), que comprueba colisiones en pasos pequeños. El prefijo _ expresa uso interno en Python; la protección funcional se obtiene aplicando estos métodos de validación.")
p("Personaje, EstadoJuego y ObservadorPuntaje son clases abstractas definidas con ABC y abstractmethod [3]. Cada una expresa un contrato distinto: comportamiento de personajes, flujo de partida y recepción de cambios de puntuación.")
h("Composición y dependencia",2)
p("El prisionero crea y posee HabilidadCarrera. Esta clase gestiona duración y recarga; el prisionero utiliza su multiplicador de velocidad. La composición permite probar esos temporizadores por separado. Por su parte, Camara.detectar(prisionero, muros) y PuertaSalida.puede_salir(prisionero) consultan parámetros sin asumir su propiedad.")
p("La vista centraliza el dibujo y no participa en las decisiones de victoria o derrota. Esta separación facilita comprobar las reglas sin una ventana y cambiar la representación visual sin modificar el conteo de puntos.")

pagina(); h("7 Implementación y trazabilidad")
tabla(["Archivo","Responsabilidad principal"],[
    ("main.py","Ventana, reloj, teclado, eventos y bucle principal."),
    ("src/juego.py","Coordinación, reinicio y delegación al estado."),
    ("src/nivel.py","Mapa, cámara, llave, salida y orden de interacciones."),
    ("src/personajes.py","Jerarquía, movimiento, patrulla y habilidad."),
    ("src/estados.py","State con cuatro estados concretos."),
    ("src/puntuacion.py","Observer, validación del puntaje y panel."),
    ("src/geometria.py","Colisiones, segmentos y visibilidad."),
    ("src/vista.py","Mapa, campos visuales, personajes e interfaz."),
    ("tests/test_juego.py","20 casos automatizados de reglas e integración."),
    ("tools/verificar.py","Ejecución de pruebas, eventos y capturas.")
],[5.2,11.8])
h("Secuencia de actualización",2)
p("En cada fotograma se procesan eventos, se leen controles y Juego delega al estado activo. Jugando actualiza Nivel; este llama update() en cada personaje y resuelve primero la detección, después la llave y finalmente la salida. Las transiciones detienen nuevas actualizaciones del nivel terminado. Vista dibuja el resultado y muestra el valor recibido por PanelPuntaje.")
h("Decisiones técnicas",2)
p("Se utilizó Pygame 2.6.1 [2] y se fijó esa versión en requirements.txt. Las verificaciones se realizaron con Python 3.12.14. El juego no requiere conexión a Internet durante la partida. Las coordenadas del modelo se expresan respecto al mapa; la vista añade el desplazamiento de pantalla únicamente al dibujar.")
p("Los campos de visión dibujados se recortan contra los mismos muros que intervienen en la detección. Se usa una aproximación de 61 rayos para mostrar cada cono; la detección utiliza un segmento directo entre el vigilante y el centro del jugador.")
h("Historial de Git",2)
p("La planificación se registró en cfaba6e. El prototipo y sus primeras evidencias se registraron en 8917f20. Los commits siguientes corresponden a la documentación y al empaquetado del parcial. El repositorio es local; no se ha publicado en un alojamiento remoto.")

pagina(); h("8 Pruebas y resultados")
p(f"El 8 de octubre se ejecutaron {REGISTRO['pruebas']} pruebas con unittest: {REGISTRO['fallos']} fallos y {REGISTRO['errores']} errores. Los resultados completos se encuentran en docs/evidencias/pruebas.txt y resultados.json. La tabla resume los escenarios realmente comprobados.")
nombres=["Personaje abstracto no instanciable","Colisión con los cuatro bordes","Muro interior sin atravesarlo","Velocidad diagonal normalizada","Alcance, cono y muro en la visión","Captura por cámara","Patrulla y captura por guardia","Captura por contacto desde atrás","Puntos de llave sin duplicación","Salida exige llave y premio único","Derrota prevalece sobre victoria","Pausa congela y permite reanudar","Reinicio restaura el estado inicial","Observer sin duplicados y con retiro","Puntaje rechaza entradas inválidas","Llamadas update polimórficas","Duración y recarga de carrera","Carrera acelera y respeta colisiones","Tiempo inválido y demora de ventana","Ruta completa mediante movimiento"]
tabla(["Caso","Escenario","Resultado"],[(f"T{i+1:02}",n,"Conforme") for i,n in enumerate(nombres)],[1.6,12.3,3.1])
p(f"La ruta completa utilizó {REGISTRO['fotogramas_ruta']} actualizaciones a 1/120 s, equivalentes a {REGISTRO['tiempo_simulado_segundos']:.3f} s de simulación, y finalizó con 600 puntos. Es un recorrido automatizado, no una medición del desempeño de una persona.")
p("La integración comprobó P, Escape, pérdida de foco, R y cierre. Las capturas provienen del renderizador Pygame con SDL en modo sin ventana. Falta el ensayo humano con teclado en el equipo de exposición; las pruebas automatizadas no sustituyen esa revisión de interacción y rendimiento.")

pagina(); h("9 Evidencia de los estados")
imagen(RAIZ/"docs/evidencias/02_pausa.png",5.05,"Figura 5. Pausa alcanzada mediante un evento de teclado P. Escape devuelve la partida a Jugando.")
imagen(RAIZ/"docs/evidencias/03_victoria.png",5.05,"Figura 6. Victoria obtenida mediante la ruta completa automatizada, con 600 puntos y sin teletransportes en ese recorrido.")
p("La captura de derrota se conserva en docs/evidencias/04_derrota.png. Se preparó un contacto con el guardia para comprobar ese resultado y su representación visual.")

pagina(); h("10 Manual de instalación y demostración")
h("Preparación del equipo",2)
p("Se recomienda utilizar Python 3.12 para reproducir el entorno validado. Extraer el paquete, abrir una terminal dentro de la carpeta Proyecto_Fuga_Prision y ejecutar los siguientes comandos. La primera instalación de Pygame necesita acceso a Internet.")
codigo("python -m venv .venv\n.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n.venv\\Scripts\\python.exe main.py")
p("En Windows, si el comando python no corresponde a Python 3.12, usar py -3.12 para crear el entorno. Si aparece No module named pygame, comprobar que la instalación y la ejecución utilizan el mismo intérprete del entorno virtual.")
h("Controles",2)
tabla(["Acción","Tecla"],[("Mover","WASD o flechas"),("Activar carrera","Shift; requiere estar disponible"),("Pausar o reanudar","P o Escape"),("Reiniciar tras victoria o derrota","R"),("Salir","Cerrar la ventana")],[10,7])
h("Verificación reproducible",2)
codigo(".venv\\Scripts\\python.exe -m unittest discover -s tests -v\n.venv\\Scripts\\python.exe tools\\verificar.py")
p("El primer comando ejecuta las reglas. El segundo también regenera cuatro capturas y el resumen JSON; usa SDL sin ventana. La carpeta .venv no forma parte del paquete y se debe crear en cada equipo.")
h("Secuencia de demostración",2)
p("1. Mostrar el objetivo y los elementos del mapa. 2. Moverse contra un muro. 3. Activar Shift y observar su recarga. 4. Pausar y comprobar que el guardia y el temporizador se detienen. 5. Entrar en vigilancia para mostrar derrota y reiniciar con R. 6. Recoger la llave, mostrar 100 puntos y alcanzar la salida para obtener 600.")
p("Para ensayar una victoria sencilla: subir por el pasillo izquierdo, recoger la llave, continuar hasta la fila superior transitable, desplazarse hacia la derecha por encima de la cámara y bajar a la salida. El mapa inicial prioriza la demostración de las reglas; su dificultad se ajustará después del ensayo.")

pagina(); h("11 Trabajo pendiente y conclusiones")
h("Ampliaciones para el proyecto final",2)
p("El siguiente incremento debe añadir al menos un segundo nivel para atender el requisito de distintos niveles. Se mantienen como propuesta tres escenarios: pabellón, vigilancia y salida exterior. Cada nivel deberá variar obstáculos o vigilancia y conservar el puntaje acumulado al avanzar.")
p("La API externa sigue pendiente de aclaración con el docente. Si se exige, se incorporará una clase APIClient que concentre la comunicación y maneje fallos de red sin impedir una partida local. No se atribuye al prototipo una integración que todavía no existe.")
tabla(["Fecha","Actividad prevista hasta el parcial"],[
    ("8 de octubre","Implementación y comprobación automatizada del prototipo; documentación del avance."),
    ("9 al 10 de octubre","Revisar consigna del parcial con el docente y ensayar el juego con teclado."),
    ("11 al 12 de octubre","Corregir problemas observados y ajustar dificultad; mantener pruebas."),
    ("13 al 14 de octubre","Revisar correspondencia UML y código y practicar la explicación de POO."),
    ("15 de octubre","Probar instalación en el equipo de presentación y ensayar la demostración."),
    ("16 de octubre","Entregar el parcial y presentar el avance alcanzado.")
],[3.4,13.6])
h("Conclusiones del avance",2)
p("El prototipo implementa un ciclo completo de juego y demuestra los dos patrones seleccionados. State determina qué entradas y actualizaciones admite cada momento de la partida. Observer comunica los cambios del puntaje sin introducir operaciones gráficas en la clase que los calcula.")
p("Las pruebas ejecutadas respaldan colisiones, detección, prioridades, temporizadores, reinicio y una ruta de victoria. El alcance validado corresponde a un nivel y a ejecución automatizada. Las próximas revisiones deben centrarse en la experiencia de uso y en ampliar el contenido del juego sin alterar los contratos que ya se pueden comprobar.")
h("Referencias",2)
p("[1] Programación Orientada a Objetos. Proyecto Final POO. Consigna del curso, ciclo 2026-2, 3 páginas. Archivo proporcionado por el estudiante.")
p("[2] Pygame. Documentación oficial. https://www.pygame.org/docs/ . Consultada el 8 de octubre de 2026.")
p("[3] Python Software Foundation. abc — Abstract Base Classes. https://docs.python.org/3/library/abc.html . Consultada el 8 de octubre de 2026.")

doc.core_properties.title="Fuga de la Prisión Informe de avance parcial"
doc.core_properties.subject="Proyecto de Programación Orientada a Objetos"
doc.core_properties.author="Jose Fernando Chavez Lozano; Mickele Carlos Torres Velazco; Richard Aaron Soriano Cordova; Piero Cemei Romero Meza"
doc.core_properties.keywords="Python, POO, State, Observer, videojuego 2D"
for raiz_xml in (doc._element, doc.styles.element):
    for borde in list(raiz_xml.iter(qn("w:pBdr"))):
        borde.getparent().remove(borde)
doc.save(SALIDA/"Informe_Parcial_Fuga_Prision.docx")
(RAIZ/"docs/02_informe_parcial.md").write_text("# Fuga de la Prisión\n\n"+"\n".join(md),encoding="utf-8")
print(SALIDA/"Informe_Parcial_Fuga_Prision.docx")
